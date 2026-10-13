"""World Bank historical income classification (OGHIST.xlsx) → dbt seed CSV.

The World Bank classifies economies by income each July 1 for its fiscal year (FY27 runs from
2026-07-01 to 2027-06-30) and publishes the full history since FY89 as an Excel workbook. The
Indicators API only returns the current group, so this history is the source for the SCD2
dimension in dbt.

    python -m ingestion.worldbank.income_history            # download, verify, write the seed
    python -m ingestion.worldbank.income_history --file OGHIST.xlsx

The workbook is replaced each July. Its SHA-256 is pinned: a new release stops the script, so the
change is reviewed (new fiscal year, revised history) before the pin and the seed are updated.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import sys
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import requests
from openpyxl import load_workbook

URL = "https://ddh-openapi.worldbank.org/resources/DR0095334/download"
SHA256 = "f9e3908b9ecd3f9480e6e00b43b1beb2cd04c9fbe8e04db3b4bce47d621fb027"
SHEET = "Country Analytical History"
SEED = Path("transform/dbt/seeds/wb_income_history.csv")

# Workbook codes → the income codes used by the Indicators API (silver_worldbank_countries).
INCOME_CODES = {"L": "LIC", "LM": "LMC", "LM*": "LMC", "UM": "UMC", "H": "HIC"}
# "..": no data for the economy that year; empty: not classified. Neither is a classification.
NOT_CLASSIFIED = {None, "", ".."}
COUNTRY_CODE = re.compile(r"[A-Z]{3}[a-z]?")  # e.g. YUGf, former Yugoslavia
FISCAL_YEAR = re.compile(r"FY(\d{2})")


class IncomeHistoryError(RuntimeError):
    pass


@dataclass(frozen=True)
class Classification:
    country_code: str
    country_name: str
    fiscal_year: int
    income_code: str

    @property
    def valid_from(self) -> date:
        """First day of the fiscal year: FY27 starts on 2026-07-01."""
        return date(self.fiscal_year - 1, 7, 1)


def fiscal_year(label: str) -> int:
    match = FISCAL_YEAR.fullmatch(label.strip())
    if not match:
        raise IncomeHistoryError(f"Unexpected fiscal year label {label!r}")
    yy = int(match.group(1))
    return 1900 + yy if yy >= 50 else 2000 + yy


def parse(rows: Iterable[tuple]) -> Iterator[Classification]:
    """Yield one classification per economy and fiscal year from the sheet's rows."""
    years: dict[int, int] | None = None
    for row in rows:
        first, second = (row + (None, None))[:2]
        if years is None:
            if second == "Bank's fiscal year:":
                years = {i: fiscal_year(v) for i, v in enumerate(row) if i >= 2 and v}
            continue
        if not (isinstance(first, str) and COUNTRY_CODE.fullmatch(first)):
            continue
        for i, fy in years.items():
            value = row[i] if i < len(row) else None
            value = value.strip() if isinstance(value, str) else value
            if value in NOT_CLASSIFIED:
                continue
            if value not in INCOME_CODES:
                raise IncomeHistoryError(f"{first} FY{fy}: unknown income code {value!r}")
            yield Classification(first, str(second).strip(), fy, INCOME_CODES[value])
    if years is None:
        raise IncomeHistoryError("Header row \"Bank's fiscal year:\" not found")


def read_workbook(content: bytes) -> list[Classification]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise IncomeHistoryError(f"Sheet {SHEET!r} missing; found {workbook.sheetnames}")
    rows = list(parse(workbook[SHEET].iter_rows(values_only=True)))
    keys = {(r.country_code, r.fiscal_year) for r in rows}
    if len(keys) != len(rows):
        raise IncomeHistoryError("Duplicate economy and fiscal year rows")
    return rows


def write_seed(rows: list[Classification], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(["country_code", "country_name", "fiscal_year", "valid_from",
                         "income_code"])
        for r in sorted(rows, key=lambda r: (r.country_code, r.fiscal_year)):
            writer.writerow([r.country_code, r.country_name, r.fiscal_year,
                             r.valid_from.isoformat(), r.income_code])


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m ingestion.worldbank.income_history",
                                     description=__doc__.split("\n")[0])
    parser.add_argument("--file", type=Path, help="use a local OGHIST.xlsx instead of downloading")
    parser.add_argument("--out", type=Path, default=SEED)
    args = parser.parse_args(argv)

    if args.file:
        content = args.file.read_bytes()
    else:
        response = requests.get(URL, timeout=120)
        response.raise_for_status()
        content = response.content
    digest = hashlib.sha256(content).hexdigest()
    if digest != SHA256:
        print(f"OGHIST changed: sha256 {digest}. Review the new release, then update SHA256.",
              file=sys.stderr)
        return 1

    rows = read_workbook(content)
    write_seed(rows, args.out)
    years = sorted({r.fiscal_year for r in rows})
    print(f"{len(rows):,} classifications, {len({r.country_code for r in rows})} economies, "
          f"FY{years[0] % 100:02d}–FY{years[-1] % 100:02d} → {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
